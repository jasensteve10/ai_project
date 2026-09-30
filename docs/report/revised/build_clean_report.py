"""Build the clean editable report from the reorganized draft; no model calls."""
from pathlib import Path
import argparse
import hashlib
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
SOURCE = Path(__file__).with_name('PROJECT_REPORT_CLEAN.md')
DEST = ROOT / 'output/docx/PROJECT_REPORT_CLEAN.docx'
BLACK = RGBColor(0,0,0)


def link(p,text,url):
    rel = p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'),'214D6B'); prop.append(c)
    r.append(prop); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)


def rich(p,text):
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|https?://[^\s]+)'
    for part in re.split(pattern,text):
        if not part: continue
        if part.startswith('http'):
            link(p,part,part); continue
        run=p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('*') else part)
        if part.startswith('**'):run.bold=True
        elif part.startswith('*'):run.italic=True


def set_shade(cell,color):
    e=OxmlElement('w:shd');e.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(e)


def table(doc,rows,width_cm):
    n=len(rows[0]); header=rows[0][0]
    ratios = {2:[.25,.75],3:[.28,.43,.29],4:[.34,.22,.22,.22]}[n]
    if header=='Card type':ratios=[.19,.12,.69]
    if header=='Setting':ratios=[.24,.42,.34]
    if header=='Method':ratios=[.31,.21,.25,.23]
    if header=='Live question':ratios=[.36,.38,.13,.13]
    if header.startswith('Metric for'):ratios=[.56,.22,.22]
    t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,r in zip(t.columns,ratios):c.width=Cm(width_cm*r)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for ri,rowdata in enumerate(rows):
        row=t.add_row();trpr=row._tr.get_or_add_trPr()
        nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,(cell,text) in enumerate(zip(row.cells,rowdata)):
            cell.width=Cm(width_cm*ratios[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for side,v in [('top',72),('bottom',72),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.05
            numeric = header in ['Method','Card type'] and ci>0 and (ci<3 if header=='Card type' else True)
            if header=='Card type' and ci==2:numeric=False
            if header=='Live question' and ci>1:numeric=True
            if header.startswith('Metric for') and ci>0:numeric=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if numeric else WD_ALIGN_PARAGRAPH.LEFT
            rich(p,text)
            for run in p.runs:run.font.size=Pt(9.3);run.font.color.rgb=BLACK;run.bold=ri==0 or run.bold
            if ri==0:set_shade(cell,'E7EDF1')
            elif ri%2==0:set_shade(cell,'F7F8FA')
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(1);spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)


def build(source=SOURCE, destination=DEST, evidence_cutoff='2026-09-27'):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    sec.footer_distance=Cm(.75)
    styles=doc.styles
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
        st=styles[name];st.font.name='Arial';st.font.color.rgb=BLACK
        rpr=st.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Arial')
        color=rpr.find(qn('w:color'))
        if color is not None:
            for a in list(color.attrib):
                if 'theme' in a:del color.attrib[a]
    # Remove inherited Word template rules beneath the title.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal=styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.widow_control=True
    for name,sz,after,before in [('Title',22,12,0),('Heading 1',16,10,0),('Heading 2',11.5,5,7),('Heading 3',11,4,6)]:
        st=styles[name];st.font.size=Pt(sz);st.font.bold=True
        st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True
    styles['Title'].paragraph_format.line_spacing=1.0
    styles['Caption'].font.size=Pt(8.5);styles['Caption'].font.italic=True;styles['Caption'].font.bold=False
    styles['Caption'].paragraph_format.line_spacing=1.0;styles['Caption'].paragraph_format.space_after=Pt(7)
    for name in ['List Bullet','List Number']:
        st=styles[name];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.08
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Deep Learning with Python  |  ').font.size=Pt(8)
    r=footer.add_run();r.font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    doc.core_properties.title=source.read_text().splitlines()[0].lstrip('# ').strip()
    doc.core_properties.author='Zapfack Jasen Steve';doc.core_properties.subject='Technical project summary with EDA and preliminary evaluation'
    lines=source.read_text().splitlines();i=0;pending_pagebreak=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='<!-- pagebreak -->':pending_pagebreak=True;i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[s.strip() for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
                i+=1
            table(doc,rows,17.4);continue
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if m:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=True
            run=p.add_run();picture=run.add_picture(str((source.parent/m[2]).resolve()),width=Cm(17.4))
            picture._inline.docPr.set('descr',m[1]);i+=1;continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
            p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(level-1))
            if pending_pagebreak:p.paragraph_format.page_break_before=True;pending_pagebreak=False
            rich(p,text);i+=1;continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet');rich(p,line[2:]);i+=1;continue
        if re.match(r'^\d\. ',line):
            # Explicit numbers keep separate short lists from continuing one another.
            p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.45)
            p.paragraph_format.space_after=Pt(4);rich(p,line);i+=1;continue
        p=doc.add_paragraph(style='Caption' if line.startswith('*') and not line.startswith('**') else 'Normal')
        # Preserve deliberate author/instructor line breaks, otherwise join wrapped prose.
        rich(p,line);i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|<!--|- |\d\. )',lines[i].strip()):
            p.add_run('\n' if line.endswith('  ') or line.startswith('**Author:') or line.startswith('**Instructor:') else ' ')
            line=lines[i].strip();rich(p,line);i+=1
    destination.parent.mkdir(parents=True,exist_ok=True);doc.save(destination)
    meta={'source_draft':'project-report.docx','source_draft_sha256':hashlib.sha256((ROOT/'project-report.docx').read_bytes()).hexdigest(),
          'edited_text':str(source.relative_to(ROOT)),'docx':str(destination.relative_to(ROOT)),
          'evidence_cutoff':evidence_cutoff,'report_date':'2026-09-30'}
    (source.parent/'manifest.json').write_text(json.dumps(meta,indent=2)+'\n')
    print(destination)


if __name__=='__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    parser.add_argument('--output', type=Path, default=DEST)
    parser.add_argument('--evidence-cutoff', default='2026-09-27')
    args = parser.parse_args()
    build(args.source, args.output, args.evidence_cutoff)
